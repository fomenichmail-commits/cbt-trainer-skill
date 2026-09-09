# -*- coding: utf-8 -*-
"""Build TechDay VSK 09.09.2026 presentation (16:9 PPTX)."""
import math

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ----------------------------------------------------------------------------
# Палитра (фирменная гамма ВСК)
# ----------------------------------------------------------------------------
NAVY    = RGBColor(0x0A, 0x2A, 0x4A)
PRIMARY = RGBColor(0x0B, 0x5F, 0xA5)
ACCENT  = RGBColor(0x25, 0xA6, 0xE0)
LIGHT   = RGBColor(0xF2, 0xF5, 0xF9)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
DARK    = RGBColor(0x22, 0x2B, 0x35)
GRAY    = RGBColor(0x5F, 0x6E, 0x80)
LINE    = RGBColor(0xD8, 0xE0, 0xEA)
GOLD    = RGBColor(0xE8, 0xA2, 0x0A)
PILLBG  = RGBColor(0x10, 0x3C, 0x6E)

FONT = "Segoe UI"

# ----------------------------------------------------------------------------
# Обложка (background image 1920x1080)
# ----------------------------------------------------------------------------
def make_cover(path):
    W, H = 1920, 1080
    gw, gh = 480, 270
    grad = Image.new("RGB", (gw, gh))
    gp = grad.load()
    top = (8, 32, 64)
    bot = (12, 96, 168)
    for y in range(gh):
        t = y / gh
        c = (int(top[0] + (bot[0] - top[0]) * t),
             int(top[1] + (bot[1] - top[1]) * t),
             int(top[2] + (bot[2] - top[2]) * t))
        for x in range(gw):
            gp[x, y] = c
    img = grad.resize((W, H), Image.BILINEAR)

    ov = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    d = ImageDraw.Draw(ov)

    # точечная сетка
    for x in range(0, W, 26):
        for y in range(0, H, 26):
            d.ellipse([x - 1, y - 1, x + 1, y + 1], fill=(255, 255, 255, 16))

    cx, cy = 1500, 430
    # концентрические кольца
    for r in range(150, 560, 100):
        d.ellipse([cx - r, cy - r, cx + r, cy + r], outline=(255, 255, 255, 38), width=3)
    # сектор
    d.pieslice([cx - 540, cy - 540, cx + 540, cy + 540], 205, 340, fill=(37, 166, 224, 46))
    # линии-лучи
    for a in (0, 45, 90, 135, 180, 225, 270, 315):
        rad = math.radians(a)
        d.line([cx, cy, cx + 520 * math.cos(rad), cy + 520 * math.sin(rad)],
               fill=(255, 255, 255, 20), width=2)

    # узлы-точки на кольцах + связи
    rng = [
        (150, 40), (240, 130), (330, 60), (420, 210), (500, 320),
        (200, 290), (300, 220), (460, 100), (120, 200), (380, 330),
    ]
    pts = []
    for i, (r, a) in enumerate(rng):
        rad = math.radians(a)
        pts.append((cx + r * math.cos(rad), cy + r * math.sin(rad)))
    for i in range(len(pts)):
        for j in range(i + 1, len(pts)):
            if abs(i - j) <= 2:
                d.line([pts[i], pts[j]], fill=(255, 255, 255, 60), width=2)
    for p in pts:
        d.ellipse([p[0] - 3, p[1] - 3, p[0] + 3, p[1] + 3], fill=(37, 166, 224, 200))

    # пин локации с пульсом
    d.ellipse([cx - 150, cy - 150, cx + 150, cy + 150], outline=(255, 255, 255, 30), width=2)
    d.ellipse([cx - 95, cy - 95, cx + 95, cy + 95], outline=(255, 255, 255, 42), width=2)
    d.ellipse([cx - 62, cy - 62, cx + 62, cy + 62], fill=(37, 166, 224, 70))
    d.polygon([(cx - 13, cy + 10), (cx + 13, cy + 10), (cx, cy + 52)], fill=(37, 166, 224, 255))
    d.ellipse([cx - 30, cy - 34, cx + 30, cy + 26], fill=(37, 166, 224, 255),
              outline=(255, 255, 255, 220), width=5)
    d.ellipse([cx - 11, cy - 15, cx + 11, cy + 7], fill=(255, 255, 255, 255))

    img = Image.alpha_composite(img.convert("RGBA"), ov).convert("RGB")

    # диагональный световой блик
    hl = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    hd = ImageDraw.Draw(hl)
    hd.polygon([(-200, 700), (420, -200), (780, -200), (160, 700)], fill=(255, 255, 255, 14))
    img = Image.alpha_composite(img.convert("RGBA"), hl).convert("RGB")
    img.save(path)
    return path


# ----------------------------------------------------------------------------
# Помощники pptx
# ----------------------------------------------------------------------------
def rect(slide, x, y, w, h, fill=None, line=None, radius=None, shape=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(1)
    try:
        sp.shadow.inherit = False
    except Exception:
        pass
    return sp


def text(slide, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = para.get("align", align)
        p.space_before = Pt(para.get("before", 0))
        p.space_after = Pt(para.get("after", 0))
        if para.get("line_spacing"):
            p.line_spacing = para["line_spacing"]
        for r in para["runs"]:
            run = p.add_run()
            run.text = r["t"]
            f = run.font
            f.name = r.get("font", FONT)
            f.size = Pt(r.get("size", 14))
            f.bold = r.get("bold", False)
            f.italic = r.get("italic", False)
            f.color.rgb = r.get("color", DARK)
    return tb


def card(slide, x, y, w, h, title, title_color, bullets, body_size=12.5,
         title_size=15.5, fill=WHITE, line=LINE):
    rect(slide, x, y, w, h, fill=fill, line=line, radius=0.05,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tb = slide.shapes.add_textbox(Inches(x + 0.24), Inches(y + 0.16),
                                  Inches(w - 0.48), Inches(h - 0.32))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.space_after = Pt(7)
    r = p.add_run()
    r.text = title
    r.font.name = FONT
    r.font.size = Pt(title_size)
    r.font.bold = True
    r.font.color.rgb = title_color
    for b in bullets:
        p2 = tf.add_paragraph()
        p2.space_after = Pt(6)
        p2.line_spacing = 1.08
        r1 = p2.add_run()
        r1.text = "▪  "
        r1.font.name = FONT
        r1.font.size = Pt(body_size)
        r1.font.bold = True
        r1.font.color.rgb = ACCENT
        r2 = p2.add_run()
        r2.text = b
        r2.font.name = FONT
        r2.font.size = Pt(body_size)
        r2.font.color.rgb = DARK
    return tb


def strip(slide, x, y, w, h, prefix, body, fill=PRIMARY):
    rect(slide, x, y, w, h, fill=fill, radius=0.09, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(slide, x + 0.28, y, w - 0.56, h,
         [{"runs": [{"t": prefix, "size": 13.5, "bold": True, "color": RGBColor(0xFF, 0xE9, 0xA0)},
                    {"t": body, "size": 13.5, "color": WHITE}],
           "line_spacing": 1.0}],
         anchor=MSO_ANCHOR.MIDDLE)


def pill(slide, x, y, label, size=11, fill=WHITE, txtcolor=PRIMARY, line=PRIMARY):
    w = 0.34 + len(label) * 0.117
    rect(slide, x, y, w, 0.42, fill=fill, line=line, radius=0.5)
    text(slide, x, y, w, 0.42,
         [{"runs": [{"t": label, "size": size, "bold": True, "color": txtcolor}],
           "align": PP_ALIGN.CENTER}],
         anchor=MSO_ANCHOR.MIDDLE, wrap=False)
    return x + w + 0.14


def base_slide(prs, title_text, num, badge=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, 13.333, 7.5, fill=LIGHT)
    rect(s, 0, 0, 13.333, 0.12, fill=PRIMARY)
    text(s, 0.6, 0.42, 10.6, 0.7,
         [{"runs": [{"t": title_text, "size": 28, "bold": True, "color": NAVY}]}])
    rect(s, 0.62, 1.14, 1.7, 0.045, fill=ACCENT)
    if badge:
        w = 0.34 + len(badge) * 0.12
        bx = 13.333 - 0.6 - w
        rect(s, bx, 0.5, w, 0.42, fill=PRIMARY, radius=0.5)
        text(s, bx, 0.5, w, 0.42,
             [{"runs": [{"t": badge, "size": 11, "bold": True, "color": WHITE}],
               "align": PP_ALIGN.CENTER}],
             anchor=MSO_ANCHOR.MIDDLE, wrap=False)
    rect(s, 0.6, 7.06, 12.13, 0.012, fill=LINE)
    text(s, 0.6, 7.12, 8, 0.3,
         [{"runs": [{"t": "TechDay ВСК · 09.09.2026 · ИИ-проекты", "size": 9, "color": GRAY}]}])
    text(s, 11.4, 7.12, 1.33, 0.3,
         [{"runs": [{"t": f"{num:02d} / 10", "size": 9, "color": GRAY}],
           "align": PP_ALIGN.RIGHT}])
    return s


# ----------------------------------------------------------------------------
# Презентация
# ----------------------------------------------------------------------------
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
prs.core_properties.title = "TechDay ВСК 09.09.2026 — ИИ-проекты"
prs.core_properties.author = "ВСК"

# --- Слайд 1: обложка ------------------------------------------------------
cover = "cover.png"
make_cover(cover)
s1 = prs.slides.add_slide(prs.slide_layouts[6])
s1.shapes.add_picture(cover, 0, 0, Inches(13.333), Inches(7.5))

text(s1, 0.9, 1.45, 9, 0.4,
     [{"runs": [{"t": "САО «ВСК» · Внутренний TechDay", "size": 15, "bold": True,
                 "color": RGBColor(0xBF, 0xE3, 0xF7)}]}])
text(s1, 0.9, 2.0, 10, 1.2,
     [{"runs": [{"t": "TechDay ВСК", "size": 60, "bold": True, "color": WHITE}]}])
text(s1, 0.9, 3.15, 6, 0.7,
     [{"runs": [{"t": "09.09.2026", "size": 34, "bold": True, "color": ACCENT}]}])
rect(s1, 0.92, 4.0, 3.0, 0.05, fill=ACCENT)
text(s1, 0.9, 4.2, 9.5, 0.6,
     [{"runs": [{"t": "ИИ как способ Discovery для Product Owner", "size": 23,
                 "bold": True, "color": WHITE}]}])

px = 0.9
for label in ("GeoCum", "GEOAIN", "Веб-форма страхования грузов"):
    w = 0.5 + len(label) * 0.145
    rect(s1, px, 5.0, w, 0.5, fill=PILLBG, line=RGBColor(0x25, 0xA6, 0xE0), radius=0.5)
    text(s1, px, 5.0, w, 0.5,
         [{"runs": [{"t": label, "size": 14, "bold": True, "color": WHITE}],
           "align": PP_ALIGN.CENTER}],
         anchor=MSO_ANCHOR.MIDDLE, wrap=False)
    px += w + 0.2

text(s1, 0.9, 6.65, 11, 0.4,
     [{"runs": [{"t": "3 ИИ-проекта · проблемы → решение · стек, CI/CD и интеграции",
                 "size": 13, "color": RGBColor(0xBF, 0xE3, 0xF7)}]}])

# --- Слайд 2: Discovery ----------------------------------------------------
s2 = base_slide(prs, "ИИ как способ Discovery для Product Owner", 2)
text(s2, 0.6, 1.32, 12.13, 0.9,
     [{"runs": [{"t": "Быстрое моделирование задачи и мгновенная обратная связь от заказчика позволяют "
                      "не ошибаться с БТ/ФТ: выявляем истинные боли, формулируем корректные требования, "
                      "оцениваем сложность и трудоёмкость и подтверждаем гипотезу через MVP.",
                 "size": 14, "color": GRAY}], "line_spacing": 1.15}])

intro = [
    ("01", "Моделирование задачи",
     "Собираем описание проекта и прототип за минуты, сразу показываем заказчику и обсуждаем."),
    ("02", "Корректные БТ/ФТ",
     "Выявляем истинные боли и формулируем бизнес- и функциональные требования без неопределённости."),
    ("03", "Оценка сложности",
     "Описываем проект и заранее оцениваем сложность и трудоёмкость разработки."),
    ("04", "MVP",
     "Подтверждаем гипотезу работающим прототипом до вложений в продакшен."),
]
for i, (num, title, body) in enumerate(intro):
    x = 0.6 + i * (2.9 + 0.15)
    rect(s2, x, 1.95, 2.9, 4.1, fill=WHITE, line=LINE, radius=0.05,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s2, x, 1.95, 2.9, 0.07, fill=ACCENT)
    text(s2, x + 0.24, 2.2, 2.42, 0.6,
         [{"runs": [{"t": num, "size": 26, "bold": True, "color": ACCENT}]}])
    text(s2, x + 0.24, 2.85, 2.42, 1.2,
         [{"runs": [{"t": title, "size": 15.5, "bold": True, "color": NAVY}],
           "line_spacing": 1.05}])
    text(s2, x + 0.24, 3.85, 2.42, 1.9,
         [{"runs": [{"t": body, "size": 12.5, "color": DARK}], "line_spacing": 1.15}])
strip(s2, 0.6, 6.3, 12.13, 0.55, "Как дальше: ",
      "3 проекта — в каждом 2 слайда: проблемы → решение.")

# --- Слайды 3–4: GeoCum ----------------------------------------------------
s3 = base_slide(prs, "Проект 1 · GeoCum — кумуляция рисков", 3, "ПРОЕКТ 1 / 3")
card(s3, 0.6, 1.55, 5.95, 4.4, "Что решали — проблемы", PRIMARY, [
    "При заключении договоров страхования юридических лиц нужно рассчитывать кумуляцию рисков.",
    "Кумуляция — возможный совокупный ущерб по застрахованным объектам в одном инциденте.",
    "Необходима для корректного принятия рисков: предельная кумуляция определяется объёмом облигатора.",
])
card(s3, 6.78, 1.55, 5.95, 4.4, "Решение", NAVY, [
    "Пайплайн: адрес ЮЛ → координаты и кадастровый номер (Дадата) → контур здания (НСПД / Росреестр) → расчёт кумуляции.",
    "Кумуляция PD/BI по радиусам 25/50/75/100 м; режимы gross/net; конвертация валют (ЦБ РФ).",
    "Веб-демо: ввод адреса → зона риска на карте, рекомендуемый радиус.",
    "Отчёт XLSX + журнал действий для аудита.",
])
strip(s3, 0.6, 6.15, 12.13, 0.62, "Итог: ",
      "работающий MVP за недели — гипотеза подтверждена, андеррайтеры получают данные для принятия решений.")

s4 = base_slide(prs, "GeoCum · Задачи со звёздочкой", 4, "ПРОЕКТ 1 / 3")
card(s4, 0.6, 1.6, 5.95, 4.0, "★  Кумуляция от контура здания", GOLD, [
    "Зона риска строится не от точки, а от реального контура здания (НСПД / Росреестр).",
    "Учитывается фактическая геометрия объекта — точность оценки совокупного ущерба выше.",
])
card(s4, 6.78, 1.6, 5.95, 4.0, "★★  «Плоские» объекты и сооружения", GOLD, [
    "Расчёт распространяется на сооружения различного типа: площадки, склады, инфраструктура.",
    "Учёт протяжённой геометрии объектов вместо условной точки.",
])
strip(s4, 0.6, 5.85, 12.13, 0.62, "Статус: ",
      "кумуляция от контура реализована в MVP; «плоские» объекты — следующий шаг к продакшену.")

# --- Слайды 5–6: GEOAIN ----------------------------------------------------
s5 = base_slide(prs, "Проект 2 · GEOAIN — анализ стихийных явлений", 5, "ПРОЕКТ 2 / 3")
card(s5, 0.6, 1.55, 5.95, 4.4, "Что решали — проблемы", PRIMARY, [
    "Аналог уже существует у китайских страховых компаний — нужен собственный сервис.",
    "Стихийные явления (осадки, ветер, заморозки, жара) → массовые убытки по портфелю.",
    "По разнородным данным нужно прогнозировать тип события и площадь зоны.",
])
card(s5, 6.78, 1.55, 5.95, 4.4, "Решение", NAVY, [
    "Анализ метеоданных за последние 72 часа (Open-Meteo).",
    "Детекция явлений по порогам → объединение в полигоны зон риска.",
    "Пространственный join: компании в зоне и рядом → уровень угрозы.",
    "Каталог явлений по этапам: MVP → пожары/паводки → землетрясения/оползни.",
])
strip(s5, 0.6, 6.15, 12.13, 0.62, "Итог: ",
      "сквозной MVP — от данных до карты зон и оценок по компаниям — за короткий спринт.")

s6 = base_slide(prs, "GEOAIN · Задачи со звёздочкой", 6, "ПРОЕКТ 2 / 3")
card(s6, 0.6, 1.6, 5.95, 4.0, "★  Компании в зоне события", GOLD, [
    "Определяем перечень компаний в зоне события.",
    "Выделяем застрахованных в ВСК — мгновенная оценка экспозиции портфеля.",
])
card(s6, 6.78, 1.6, 5.95, 4.0, "★★  Прогноз на 72 часа и рекомендации", GOLD, [
    "ML-прогноз (LSTM / AR) на 72 часа вперёд.",
    "Адресные рекомендации компаниям: как защитить активы до наступления события.",
])
strip(s6, 0.6, 5.85, 12.13, 0.62, "Что дальше: ",
      "спутниковые данные (NASA FIRMS), паводок (GloFAS), интерфейс на Яндекс.Картах.")

# --- Слайды 7–8: Веб-форма страхования грузов -------------------------------
s7 = base_slide(prs, "Проект 3 · Веб-форма страхования грузов", 7, "ПРОЕКТ 3 / 3")
card(s7, 0.6, 1.55, 5.95, 4.4, "Что решали — проблемы", PRIMARY, [
    "«Заявление на страхование грузов по Генеральному полису» — 18 разделов на бумаге/PDF.",
    "Часы на заполнение, ошибки и расхождения у клиента и агента.",
    "Данные разбросаны: страхователь, грузы, маршрут, перевозчики, охрана.",
])
card(s7, 6.78, 1.55, 5.95, 4.4, "Решение", NAVY, [
    "Цифровая веб-форма в фирменном стиле ВСК (дизайн-система, синяя гамма).",
    "Два макета: «одна страница» с липким оглавлением и «пошаговый мастер» (6 шагов).",
    "Логика раскрытия подблоков и валидация; заготовки под справочники.",
])
strip(s7, 0.6, 6.15, 12.13, 0.62, "Итог: ",
      "заказчик видит и подтверждает макет до написания кода — Discovery на практике.")

s8 = base_slide(prs, "Веб-форма · Следующие шаги", 8, "ПРОЕКТ 3 / 3")
steps = [
    ("Справочники",
     "Автоподстановка контрагентов по ИНН, виды грузов, география, документы-основания."),
    ("Черновик и PDF",
     "Сохранение черновика и экспорт PDF из заполненной формы."),
    ("Генерация документа",
     "ИИ формирует заявление в Word по шаблону из ответов формы."),
]
for i, (t, b) in enumerate(steps):
    x = 0.6 + i * (3.85 + 0.19)
    rect(s8, x, 1.6, 3.85, 4.0, fill=WHITE, line=LINE, radius=0.05,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s8, x, 1.6, 3.85, 0.07, fill=ACCENT)
    text(s8, x + 0.24, 1.95, 3.37, 0.8,
         [{"runs": [{"t": t, "size": 16, "bold": True, "color": NAVY}], "line_spacing": 1.05}])
    text(s8, x + 0.24, 2.8, 3.37, 2.4,
         [{"runs": [{"t": b, "size": 12.5, "color": DARK}], "line_spacing": 1.15}])
strip(s8, 0.6, 5.85, 12.13, 0.62, "Принцип: ",
      "видимый макет → подтверждение заказчика → разработка.")

# --- Слайд 9: Стек, CI/CD, интеграции --------------------------------------
s9 = base_slide(prs, "Стек, CI/CD и интеграции", 9)

text(s9, 0.6, 1.12, 6, 0.3,
     [{"runs": [{"t": "СТЕК", "size": 12, "bold": True, "color": ACCENT}]}])
px = 0.6
for label in ("Python", "FastAPI / Flask", "Angular / JS",
              "PostgreSQL / PostGIS", "PyTorch", "GeoJSON / Leaflet"):
    px = pill(s9, px, 1.45, label)

text(s9, 0.6, 2.05, 6, 0.3,
     [{"runs": [{"t": "CI/CD", "size": 12, "bold": True, "color": ACCENT}]}])
px = 0.6
for label in ("GitLab CI", "GitHub Actions", "Entire.io"):
    px = pill(s9, px, 2.38, label)
text(s9, 0.6, 2.92, 12.13, 0.32,
     [{"runs": [{"t": "Весь цикл автоматизирован: задача → код → ревью → тесты → релиз; "
                      "Entire.io — ИИ-агент, который ведёт разработку автономно.",
                 "size": 11.5, "color": GRAY}]}])

text(s9, 0.6, 3.34, 8, 0.3,
     [{"runs": [{"t": "ИНТЕГРАЦИИ В ПРОЕКТАХ", "size": 12, "bold": True, "color": ACCENT}]}])
intg = [
    ("GeoCum", [
        "Дадата — геокодирование адресов",
        "НСПД / Росреестр (pynspd) — контуры",
        "ЦБ РФ — курсы валют",
        "PostGIS — пространственные данные",
        "Leaflet — веб-демо",
    ]),
    ("GEOAIN", [
        "Open-Meteo — метеоданные / DEM",
        "NASA FIRMS — пожары (спутник)",
        "GloFAS — паводок, речной сток",
        "Яндекс.Карты — интерфейс",
        "PyTorch — LSTM-прогноз",
    ]),
    ("Веб-форма", [
        "Справочники — следующий этап:",
        "Контрагенты · Виды грузов",
        "География · Документы-основания",
        "Сотрудники ВСК",
        "Экспорт PDF / Word по шаблону",
    ]),
]
for i, (name, items) in enumerate(intg):
    x = 0.6 + i * (3.85 + 0.19)
    rect(s9, x, 3.7, 3.85, 3.0, fill=WHITE, line=LINE, radius=0.05,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s9, x + 0.24, 3.95, 3.37, 0.5,
         [{"runs": [{"t": name, "size": 15, "bold": True, "color": NAVY}]}])
    paras = []
    for it in items:
        paras.append({"runs": [
            {"t": "▪  ", "size": 11.5, "bold": True, "color": ACCENT},
            {"t": it, "size": 11.5, "color": DARK}],
            "space_after": 5, "line_spacing": 1.08})
    text(s9, x + 0.24, 4.45, 3.37, 2.1, paras)

# --- Слайд 10: ИИ для ПО и промышленной разработки --------------------------
s10 = base_slide(prs, "ИИ для ПО и промышленной разработки", 10)
text(s10, 0.6, 1.32, 12.13, 0.5,
     [{"runs": [{"t": "Где ИИ даёт максимум — от идеи до эксплуатации.",
                 "size": 14, "color": GRAY}]}])
four = [
    ("01 · Discovery",
     "Моделирование задачи, корректные БТ/ФТ, оценка сложности и трудоёмкости, подтверждение гипотезы через MVP."),
    ("02 · Код и качество",
     "Генерация кода, код-ревью, unit-тесты, CI/CD, документация — быстрее, без потери качества."),
    ("03 · Данные и ML",
     "Разнородные данные (метео, гео, реестры, документы), обучение моделей, прогнозы и рекомендации."),
    ("04 · Эксплуатация",
     "Мониторинг, разбор инцидентов, генерация отчётов и метрик для управления продуктом."),
]
for i, (t, b) in enumerate(four):
    x = 0.6 + (i % 2) * (5.98 + 0.18)
    y = 1.95 + (i // 2) * (2.05 + 0.16)
    rect(s10, x, y, 5.98, 2.05, fill=WHITE, line=LINE, radius=0.05,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s10, x + 0.26, y + 0.2, 5.5, 0.5,
         [{"runs": [{"t": t, "size": 16.5, "bold": True, "color": NAVY}]}])
    text(s10, x + 0.26, y + 0.78, 5.5, 1.1,
         [{"runs": [{"t": b, "size": 12.5, "color": DARK}], "line_spacing": 1.15}])
strip(s10, 0.6, 6.42, 12.13, 0.55, "Вывод: ",
      "ИИ — инструмент Product Owner: быстрее понять заказчика, дешевле проверить гипотезу, "
      "быстрее довести до продакшена.")

out = "TechDay_VSK_09.09.2026.pptx"
prs.save(out)
print("OK ->", out)